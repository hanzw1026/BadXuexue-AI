# documents_assistant.py
import os
import sys
import shutil
import time
from PySide6.QtWidgets import (QWidget, QFileDialog, QMessageBox,
                               QVBoxLayout, QHBoxLayout, QPushButton,
                               QTextEdit, QComboBox, QListWidget,
                               QListWidgetItem, QLabel, QSplitter)
from PySide6.QtCore import Qt, QThread, Signal, QUrl
from PySide6.QtGui import QFont
from PySide6.QtWebEngineWidgets import QWebEngineView

import init_config
from documentsAssistantWidget import Ui_DocumentAssistantWindow
from docs_setting_dialog import DocsSettingDialog
from PySide6.QtCore import QThread, Signal
from init_config import load_doc_assistant_config
import platform


def get_platform_default_font_size():
    system = platform.system()
    if system == "Windows":
        return 13
    elif system == "Darwin":
        return 16
    else:
        return 14


DEFAULT_FONT_SIZE = get_platform_default_font_size()


# 跨平台路径处理
def get_app_dir():
    """获取应用程序所在目录（跨平台）"""
    if getattr(sys, 'frozen', False):
        # 打包后的exe运行
        return os.path.dirname(sys.executable)
    else:
        # 开发环境
        return os.path.dirname(os.path.abspath(__file__))


def ensure_empty_templates():
    """确保 docs_template 文件夹中有空模板文件"""
    app_dir = get_app_dir()
    template_dir = os.path.join(app_dir, "docs_template")
    os.makedirs(template_dir, exist_ok=True)

    empty_templates = {
        "_请上传模板或选择模板_.docx": " ",
        "_请上传模板或选择模板_.xlsx": " ",
        "_请上传模板或选择模板_.pptx": " "
    }

    for filename, content in empty_templates.items():
        filepath = os.path.join(template_dir, filename)
        if not os.path.exists(filepath):
            if filename.endswith('.docx'):
                from docx import Document
                doc = Document()
                doc.save(filepath)
            elif filename.endswith('.xlsx'):
                from openpyxl import Workbook
                wb = Workbook()
                wb.save(filepath)
            elif filename.endswith('.pptx'):
                from pptx import Presentation
                prs = Presentation()
                prs.save(filepath)
            else:
                with open(filepath, 'w', encoding='utf-8') as f:
                    f.write(content)

            if init_config.isInTestMode:
                print(f"📄 已创建空模板：{filename}")


class DocumentAssistantWindow(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.ui = Ui_DocumentAssistantWindow()
        self.ui.setupUi(self)

        # 设置整个窗口的默认字体
        font = QFont()
        font.setPointSize(DEFAULT_FONT_SIZE)
        self.setFont(font)

        self.setWindowFlags(Qt.Window)

        # 跨平台目录设置
        app_dir = get_app_dir()
        self.template_dir = os.path.join(app_dir, "docs_template")
        self.output_dir = os.path.join(app_dir, "outputs")

        # 创建目录
        os.makedirs(self.template_dir, exist_ok=True)
        os.makedirs(self.output_dir, exist_ok=True)

        # 确保空模板存在
        ensure_empty_templates()

        self.isInDarkMode = True

        # 初始化界面
        self.init_ui()
        self.init_signals()

        self.ui.combo_format.setStyleSheet(f"font-size: {DEFAULT_FONT_SIZE}pt;")
        self.ui.btn_upload_template.setStyleSheet(f"font-size: {DEFAULT_FONT_SIZE}pt;")
        self.ui.btn_delete_template.setStyleSheet(f"font-size: {DEFAULT_FONT_SIZE}pt;")
        self.ui.btn_generate.setStyleSheet(f"font-size: {DEFAULT_FONT_SIZE}pt;")
        self.ui.btn_clear.setStyleSheet(f"font-size: {DEFAULT_FONT_SIZE}pt;")
        self.ui.btn_settings.setStyleSheet(f"font-size: {DEFAULT_FONT_SIZE}pt;")

        # 加载数据
        self.load_templates()

    def init_ui(self):
        """初始化界面组件"""
        # ===== 根据暗黑模式设置预览区默认提示 =====
        if self.isInDarkMode:
            # 暗黑模式：深色背景，灰色文字
            placeholder_style = """
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {
                        background-color: #1e1e1e;
                        color: #888;
                        font-family: '微软雅黑', Arial, sans-serif;
                        display: flex;
                        justify-content: center;
                        align-items: center;
                        height: 100vh;
                        margin: 0;
                    }
                    .placeholder {
                        text-align: center;
                        font-size: {DEFAULT_FONT_SIZE - 2}pt;
                    }
                </style>
            </head>
            <body>
                <div class="placeholder">📁 选择模板后预览</div>
            </body>
            </html>
            """
            self.ui.web_template_preview.setHtml(placeholder_style)

            output_placeholder = placeholder_style.replace("📁 选择模板后预览", "✨ 生成文档后预览")
            self.ui.web_output_preview.setHtml(output_placeholder)
        else:
            # 亮色模式：白色背景，灰色文字
            self.ui.web_template_preview.setHtml(
                "<p style='color:gray;text-align:center;padding:20px'>📁 选择模板后预览</p>"
            )
            self.ui.web_output_preview.setHtml(
                "<p style='color:gray;text-align:center;padding:20px'>✨ 生成文档后预览</p>"
            )

        # 添加格式选项
        self.ui.combo_format.clear()
        self.ui.combo_format.addItems(["Word (.docx)", "Excel (.xlsx)", "PPT (.pptx)"])

        # ===== 根据暗黑模式设置输入框样式 =====
        if self.isInDarkMode:
            if platform.system() == "Windows":
                self.ui.edit_prompt.setStyleSheet("""
                    QTextEdit {
                        background-color: #3c3c3c;
                        color: #FFFFFB;
                        font-size: 13pt;
                        border: 1px solid #555;
                        border-radius: 4px;
                        padding: 8px;
                        font-family: "Microsoft YaHei";
                    }
                """)
                self.ui.edit_prompt.setPlaceholderText("输入你的需求描述...")
            elif platform.system() == "Darwin":
                self.ui.edit_prompt.setStyleSheet("""
                                    QTextEdit {
                                        background-color: #3c3c3c;
                                        color: #FFFFFB;
                                        font-size: 16pt;
                                        border: 1px solid #555;
                                        border-radius: 4px;
                                        padding: 8px;
                                    }
                                """)
                self.ui.edit_prompt.setPlaceholderText("输入你的需求描述...")
            else:
                self.ui.edit_prompt.setStyleSheet("""
                                                    QTextEdit {
                                                        background-color: #3c3c3c;
                                                        color: #FFFFFB;
                                                        border: 1px solid #555;
                                                        border-radius: 4px;
                                                        padding: 8px;
                                                    }
                                                """)
                self.ui.edit_prompt.setPlaceholderText("输入你的需求描述...")

        else:
            # 亮色模式可以保持默认或设置浅色样式
            if platform.system() == "Windows":
                self.ui.edit_prompt.setStyleSheet("""
                    QTextEdit {
                        background-color: #ffffff;
                        color: #333333;
                        font-size: 13pt;
                        font-family: "Microsoft YaHei";
                        border: 1px solid #ccc;
                        border-radius: 4px;
                        padding: 8px;
                    }
                """)
                self.ui.edit_prompt.setPlaceholderText("输入你的需求描述...")
            elif platform.system() == "Darwin":
                self.ui.edit_prompt.setStyleSheet("""
                                    QTextEdit {
                                        background-color: #ffffff;
                                        color: #333333;
                                        font-size: 16pt;
                                        border: 1px solid #ccc;
                                        border-radius: 4px;
                                        padding: 8px;
                                    }
                                """)
                self.ui.edit_prompt.setPlaceholderText("输入你的需求描述...")
            else:
                self.ui.edit_prompt.setStyleSheet("""
                                                    QTextEdit {
                                                        background-color: #ffffff;
                                                        color: #333333;
                                                        border: 1px solid #ccc;
                                                        border-radius: 4px;
                                                        padding: 8px;
                                                    }
                                                """)
                self.ui.edit_prompt.setPlaceholderText("输入你的需求描述...")

        # ===== 动态添加缩放按钮（模板预览区）=====
        self._add_zoom_controls_to_layout(
            self.ui.verticalLayout_2,  # 模板预览区的布局
            self.ui.web_template_preview,  # 对应的WebView
            "模板预览"  # 标题
        )

        # ===== 动态添加缩放按钮（输出预览区）=====
        self._add_zoom_controls_to_layout(
            self.ui.verticalLayout_3,  # 输出预览区的布局
            self.ui.web_output_preview,  # 对应的WebView
            "输出预览"  # 标题
        )

    def _add_zoom_controls_to_layout(self, layout, web_view, title):
        """给指定布局添加缩放控件（暗黑模式适配）"""
        from PySide6.QtWidgets import QHBoxLayout, QLabel, QPushButton, QSpacerItem, QSizePolicy

        zoom_layout = QHBoxLayout()
        zoom_layout.setContentsMargins(0, 0, 0, 5)

        # 标题标签（根据暗黑模式设置颜色）
        title_label = QLabel(title)
        if self.isInDarkMode:
            title_label.setStyleSheet("color: #e0e0e0; font-weight: bold; font-size: {DEFAULT_FONT_SIZE}pt;")
        else:
            title_label.setStyleSheet("color: #333; font-weight: bold; font-size: {DEFAULT_FONT_SIZE}pt;")
        zoom_layout.addWidget(title_label)

        zoom_layout.addStretch()

        # 按钮样式（暗黑模式适配）
        if self.isInDarkMode:
            btn_style = """
                QPushButton {
                    background-color: #555555;
                    color: white;
                    font-size: {DEFAULT_FONT_SIZE}pt;
                    font-weight: bold;
                    border: none;
                    border-radius: 4px;
                    padding: 4px 8px;
                    min-width: 30px;
                    min-height: 25px;
                }
                QPushButton:hover {
                    background-color: #777777;
                }
                QPushButton:pressed {
                    background-color: #444444;
                }
            """
            label_style = "color: #e0e0e0; font-size: {DEFAULT_FONT_SIZE}pt;"
        else:
            btn_style = """
                QPushButton {
                    background-color: #f0f0f0;
                    color: #333;
                    font-size: {DEFAULT_FONT_SIZE}pt;
                    font-weight: bold;
                    border: 1px solid #ccc;
                    border-radius: 4px;
                    padding: 4px 8px;
                    min-width: 30px;
                    min-height: 25px;
                }
                QPushButton:hover {
                    background-color: #e0e0e0;
                }
            """
            label_style = "color: #666; font-size: {DEFAULT_FONT_SIZE}pt;"

        # 缩小按钮
        btn_zoom_out = QPushButton("−")
        btn_zoom_out.setFixedSize(30, 25)
        btn_zoom_out.setToolTip("缩小 (Ctrl+滚轮)")
        btn_zoom_out.setStyleSheet(btn_style)

        # 放大按钮
        btn_zoom_in = QPushButton("+")
        btn_zoom_in.setFixedSize(30, 25)
        btn_zoom_in.setToolTip("放大 (Ctrl+滚轮)")
        btn_zoom_in.setStyleSheet(btn_style)

        # 重置按钮
        btn_zoom_reset = QPushButton("重置")
        btn_zoom_reset.setFixedSize(50, 25)
        btn_zoom_reset.setToolTip("重置缩放")
        btn_zoom_reset.setStyleSheet(btn_style)

        # 缩放比例显示
        zoom_label = QLabel("100%")
        zoom_label.setFixedWidth(45)
        zoom_label.setAlignment(Qt.AlignCenter)
        zoom_label.setStyleSheet(label_style)

        # 保存引用
        web_view.zoom_label = zoom_label

        # 绑定事件
        btn_zoom_in.clicked.connect(lambda: self._zoom_webview(web_view, 0.1, zoom_label))
        btn_zoom_out.clicked.connect(lambda: self._zoom_webview(web_view, -0.1, zoom_label))
        btn_zoom_reset.clicked.connect(lambda: self._reset_zoom(web_view, zoom_label))

        zoom_layout.addWidget(btn_zoom_out)
        zoom_layout.addWidget(btn_zoom_in)
        zoom_layout.addWidget(btn_zoom_reset)
        zoom_layout.addWidget(zoom_label)

        # 插入到布局的第1个位置
        layout.insertLayout(1, zoom_layout)

    def _zoom_webview(self, web_view, delta, zoom_label):
        """缩放WebView"""
        current = web_view.zoomFactor()
        new_zoom = max(0.5, min(2.0, current + delta))
        web_view.setZoomFactor(new_zoom)
        zoom_label.setText(f"{int(new_zoom * 100)}%")

    def _reset_zoom(self, web_view, zoom_label):
        """重置WebView缩放"""
        web_view.setZoomFactor(1.0)
        zoom_label.setText("100%")

    def init_signals(self):
        """绑定信号槽"""
        self.ui.btn_upload_template.clicked.connect(self.upload_template)
        self.ui.btn_delete_template.clicked.connect(self.delete_template)
        self.ui.btn_clear.clicked.connect(self.clear_prompt)
        self.ui.btn_generate.clicked.connect(self.generate_document)
        self.ui.btn_settings.clicked.connect(self.open_settings)
        self.ui.list_templates.itemClicked.connect(self.on_template_selected)
        # 右键菜单刷新
        self.ui.list_templates.setContextMenuPolicy(Qt.CustomContextMenu)
        self.ui.list_templates.customContextMenuRequested.connect(self.show_template_context_menu)

    def show_template_context_menu(self, position):
        """显示模板列表右键菜单"""
        from PySide6.QtWidgets import QMenu

        menu = QMenu()

        # 刷新列表
        refresh_action = menu.addAction("🔄 刷新列表")
        refresh_action.triggered.connect(self.load_templates)

        menu.addSeparator()

        # 取消选择（选中空模板）
        cancel_action = menu.addAction("🔘 取消选择模板")
        cancel_action.triggered.connect(self.select_empty_template)

        # 获取当前右键点击的项
        current_item = self.ui.list_templates.itemAt(position)
        if current_item:
            menu.addSeparator()
            delete_action = menu.addAction("🗑️ 删除此模板")
            delete_action.triggered.connect(self.delete_template)

        menu.exec_(self.ui.list_templates.mapToGlobal(position))

    def select_empty_template(self):
        """选中空模板（取消模板选择）"""
        # 查找空模板
        for i in range(self.ui.list_templates.count()):
            item = self.ui.list_templates.item(i)
            if "请上传模板或选择模板" in item.text():
                self.ui.list_templates.setCurrentItem(item)
                # 触发预览
                self.on_template_selected(item)
                if init_config.isInTestMode:
                    print(f"✅ 已切换到空模板：{item.text()}")
                return

        # 如果没找到空模板，清除选中
        self.ui.list_templates.clearSelection()
        self.ui.list_templates.setCurrentItem(None)
        self._reset_template_preview()
        if init_config.isInTestMode:
            print(f"❌ 未找到空模板，已清除选中")

    def load_templates(self):
        """加载模板列表（每次调用都会重新扫描文件夹）"""
        self.ui.list_templates.clear()

        if not os.path.exists(self.template_dir):
            os.makedirs(self.template_dir, exist_ok=True)

        extensions = ('.docx', '.xlsx', '.pptx', '.dotx', '.xltx')
        templates = []

        for f in sorted(os.listdir(self.template_dir)):
            if f.lower().endswith(extensions):
                templates.append(f)
                self.ui.list_templates.addItem(f)

        # 如果没有模板，显示提示
        if not templates:
            self.ui.web_template_preview.setHtml(
                "<p style='color:gray;text-align:center;padding:20px'>"
                "📁 暂无模板<br>"
                "请点击「上传文档模板」添加，或直接将文件放入 docs_template 文件夹"
                "</p>"
            )

    def upload_template(self):
        """上传模板文件（跨平台文件对话框）"""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "选择模板文件",
            "",
            "模板文件 (*.docx *.xlsx *.pptx *.dotx *.xltx)"
        )

        if not file_path:
            return

        file_name = os.path.basename(file_path)
        dest_path = os.path.join(self.template_dir, file_name)

        # 检查文件是否已存在
        if os.path.exists(dest_path):
            reply = QMessageBox.question(
                self,
                "文件已存在",
                f"模板「{file_name}」已存在，是否覆盖？",
                QMessageBox.Yes | QMessageBox.No
            )
            if reply == QMessageBox.No:
                return

        try:
            shutil.copy2(file_path, dest_path)
            self.load_templates()
            QMessageBox.information(self, "成功", f"模板「{file_name}」已添加")
        except Exception as e:
            QMessageBox.critical(self, "错误", f"上传失败：{str(e)}")

    def delete_template(self):
        """删除选中的模板"""
        current_item = self.ui.list_templates.currentItem()
        if not current_item:
            QMessageBox.warning(self, "提示", "请先选中要删除的模板")
            return

        template_name = current_item.text()
        template_path = os.path.join(self.template_dir, template_name)

        reply = QMessageBox.question(
            self,
            "确认删除",
            f"确定要删除模板「{template_name}」吗？",
            QMessageBox.Yes | QMessageBox.No
        )

        if reply == QMessageBox.Yes:
            try:
                os.remove(template_path)
                self.load_templates()
                self.ui.web_template_preview.setHtml(
                    "<p style='color:gray;text-align:center;padding:20px'>✅ 模板已删除</p>"
                )
                QMessageBox.information(self, "成功", f"模板「{template_name}」已删除")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"删除失败：{str(e)}")

    def on_template_selected(self, item):
        """预览选中的模板（跨平台文件读取）"""
        template_name = item.text()
        template_path = os.path.join(self.template_dir, template_name)

        if not os.path.exists(template_path):
            self.ui.web_template_preview.setHtml(
                "<p style='color:red'>文件不存在</p>"
            )
            return

        ext = os.path.splitext(template_name)[1].lower()

        try:
            if ext == '.docx':
                if self.isInDarkMode:
                    self._preview_word_dark(template_path)
                else:
                    self._preview_word(template_path)
            elif ext == '.xlsx':
                if self.isInDarkMode:
                    self._preview_excel_dark(template_path)
                else:
                    self._preview_excel(template_path)
            else:
                self.ui.web_template_preview.setHtml(
                    f"<p>📄 模板：{template_name}</p>"
                    f"<p>格式：{ext}</p>"
                    f"<p>大小：{os.path.getsize(template_path)} 字节</p>"
                    f"<p>提示：PPT预览功能开发中</p>"
                )
        except Exception as e:
            self.ui.web_template_preview.setHtml(
                f"<p style='color:red'>预览失败：{str(e)}</p>"
            )

    def _preview_word(self, file_path):
        """预览 Word 文档（保留基本格式）"""
        try:
            from docx import Document
            doc = Document(file_path)

            # 构建HTML内容
            html_parts = []

            # 处理段落
            para_count = 0
            for para in doc.paragraphs:
                if para_count >= 50:  # 限制显示50段
                    break

                if not para.text.strip():
                    continue

                # 获取段落对齐方式
                align_css = ""
                if para.alignment is not None:
                    align_map = {
                        0: "left", 1: "center", 2: "right", 3: "justify"
                    }
                    align_css = f"text-align: {align_map.get(para.alignment, 'left')};"

                # 处理段落中的文本和格式
                para_html = []
                for run in para.runs:
                    text = run.text
                    if not text:
                        continue

                    styles = []
                    if run.bold:
                        styles.append("font-weight: bold")
                    if run.italic:
                        styles.append("font-style: italic")
                    if run.underline:
                        styles.append("text-decoration: underline")

                    style_str = "; ".join(styles)
                    if style_str:
                        para_html.append(f'<span style="{style_str}">{text}</span>')
                    else:
                        para_html.append(text)

                if para_html:
                    html_parts.append(
                        f'<p style="margin:0 0 8px 0; {align_css}">{"".join(para_html)}</p>'
                    )
                    para_count += 1

            # 处理表格
            tables_html = []
            for i, table in enumerate(doc.tables[:3]):  # 最多显示3个表格
                table_html = f'<h4>表格 {i + 1}</h4><table border="1" cellpadding="5" style="border-collapse: collapse; width: 100%;">'
                for row in table.rows[:10]:  # 每个表格最多10行
                    table_html += '<tr>'
                    for cell in row.cells:
                        cell_text = cell.text[:100]  # 每个单元格最多100字
                        table_html += f'<td style="border: 1px solid #ddd; padding: 5px;">{cell_text}</td>'
                    table_html += '</tr>'
                table_html += '</table><br>'
                tables_html.append(table_html)

            # 完整HTML
            full_html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', '宋体', Arial, sans-serif;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        padding: 20px;
                        margin: 0;
                        line-height: 1.4;
                    }}
                    .info-bar {{
                        background-color: #f5f5f5;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #666;
                    }}
                    h4 {{
                        margin: 15px 0 5px 0;
                        color: #333;
                    }}
                    table {{
                        margin: 10px 0;
                    }}
                    td {{
                        vertical-align: top;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📄 文档共 {len(doc.paragraphs)} 段，{len(doc.tables)} 个表格
                    （预览前50段，前3个表格）
                    💡 提示：使用上方工具栏的 + / - 按钮可缩放预览
                </div>
                {''.join(html_parts) if html_parts else '<p style="color:gray;">无文字内容可预览</p>'}
                {''.join(tables_html) if tables_html else ''}
            </body>
            </html>
            """

            self.ui.web_template_preview.setHtml(full_html)

        except Exception as e:
            self.ui.web_template_preview.setHtml(
                f"<p style='color:red;padding:20px;'>❌ Word预览失败：{str(e)}<br><br>请确保已安装 python-docx</p>"
            )

    def _preview_excel(self, file_path):
        """预览 Excel 文档（支持合并单元格）"""
        try:
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            # 加载工作簿
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active

            # 获取合并单元格区域
            merged_ranges = list(ws.merged_cells.ranges)

            # 获取数据范围
            min_row = ws.min_row or 1
            max_row = min(ws.max_row, 50)  # 最多50行
            min_col = ws.min_column or 1
            max_col = min(ws.max_column, 20)  # 最多20列

            # 构建数据矩阵
            data = []
            for row in range(min_row, max_row + 1):
                row_data = []
                for col in range(min_col, max_col + 1):
                    cell = ws.cell(row, col)
                    value = cell.value

                    # 检查是否在合并单元格中（且不是左上角）
                    is_merged = False
                    for merged_range in merged_ranges:
                        if (merged_range.min_row <= row <= merged_range.max_row and
                                merged_range.min_col <= col <= merged_range.max_col):
                            if row == merged_range.min_row and col == merged_range.min_col:
                                # 左上角单元格，正常显示
                                break
                            else:
                                # 被合并的单元格，置空
                                is_merged = True
                                break

                    if is_merged:
                        row_data.append('')
                    else:
                        # 处理空值
                        if value is None:
                            row_data.append('')
                        elif isinstance(value, (int, float)):
                            # 保留两位小数
                            row_data.append(f'{value:.2f}' if value % 1 != 0 else str(value))
                        else:
                            row_data.append(str(value)[:100])
                data.append(row_data)

            # 获取列宽
            col_widths = {}
            for col in range(min_col, max_col + 1):
                col_letter = get_column_letter(col)
                width = ws.column_dimensions[col_letter].width
                if width:
                    col_widths[col] = min(width, 200)
                else:
                    col_widths[col] = 100

            # 构建表格HTML
            html_rows = []

            # 表头（列号）
            html_rows.append('<thead><tr>')
            for col in range(min_col, max_col + 1):
                col_letter = get_column_letter(col)
                width = col_widths.get(col, 100)
                html_rows.append(f'<th style="min-width: {width}px;">{col_letter}</th>')
            html_rows.append('</thead>')

            # 表体
            html_rows.append('<tbody>')
            for row_idx, row_data in enumerate(data):
                html_rows.append('<tr>')
                for value in row_data:
                    html_rows.append(f'<td>{value}</td>')
                html_rows.append('</tr>')
            html_rows.append('</tbody>')

            # 处理合并单元格的HTML
            # 需要生成 col/rowspan 属性
            # 这里简化处理，用CSS模拟合并效果

            # 更精确的方法：构建带合并信息的HTML
            html_table = self._build_excel_html_with_merges(ws, min_row, max_row, min_col, max_col)

            if html_table:
                final_html = html_table
            else:
                # 降级方案
                final_html = f"""
                <table class="excel-table">
                    {''.join(html_rows)}
                </table>
                """

            # 完整HTML
            full_html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', 'Segoe UI', Arial, sans-serif;
                        padding: 10px;
                        margin: 0;
                        background-color: #fff;
                    }}
                    .excel-table {{
                        border-collapse: collapse;
                        width: 100%;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                    }}
                    .excel-table th {{
                        background-color: #f2f2f2;
                        border: 1px solid #ddd;
                        padding: 8px;
                        text-align: center;
                        font-weight: bold;
                        position: sticky;
                        top: 0;
                    }}
                    .excel-table td {{
                        border: 1px solid #ddd;
                        padding: 6px 8px;
                        text-align: left;
                        white-space: nowrap;
                    }}
                    .excel-table tr:nth-child(even) {{
                        background-color: #f9f9f9;
                    }}
                    .excel-table tr:hover {{
                        background-color: #e8f4f8;
                    }}
                    .info-bar {{
                        background-color: #f5f5f5;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #666;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📊 工作表：{ws.title}，共 {ws.max_row} 行，{ws.max_column} 列
                    （预览区域：{min_row}-{max_row}行，{min_col}-{max_col}列）
                    合并单元格：{len(merged_ranges)} 处
                    💡 提示：使用上方工具栏的 + / - 按钮可缩放预览
                </div>
                {final_html}
            </body>
            </html>
            """

            self.ui.web_template_preview.setHtml(full_html)

        except Exception as e:
            self.ui.web_template_preview.setHtml(
                f"<p style='color:red;padding:20px;'>❌ Excel预览失败：{str(e)}<br><br>请确保已安装 openpyxl</p>"
            )

    def _preview_word_dark(self, file_path):
        """预览 Word 文档（保留基本格式）- 暗黑模式"""
        try:
            from docx import Document
            doc = Document(file_path)

            # 构建HTML内容
            html_parts = []

            # 处理段落
            para_count = 0
            for para in doc.paragraphs:
                if para_count >= 50:
                    break
                if not para.text.strip():
                    continue

                align_css = ""
                if para.alignment is not None:
                    align_map = {0: "left", 1: "center", 2: "right", 3: "justify"}
                    align_css = f"text-align: {align_map.get(para.alignment, 'left')};"

                para_html = []
                for run in para.runs:
                    text = run.text
                    if not text:
                        continue

                    styles = []
                    if run.bold:
                        styles.append("font-weight: bold")
                    if run.italic:
                        styles.append("font-style: italic")
                    if run.underline:
                        styles.append("text-decoration: underline")

                    style_str = "; ".join(styles)
                    if style_str:
                        para_html.append(f'<span style="{style_str}">{text}</span>')
                    else:
                        para_html.append(text)

                if para_html:
                    html_parts.append(
                        f'<p style="margin:0 0 8px 0; {align_css}">{"".join(para_html)}</p>'
                    )
                    para_count += 1

            # 处理表格
            tables_html = []
            for i, table in enumerate(doc.tables[:3]):
                table_html = f'<h4>表格 {i + 1}</h4><table border="1" cellpadding="5" style="border-collapse: collapse; width: 100%;">'
                for row in table.rows[:10]:
                    table_html += '<tr>'
                    for cell in row.cells:
                        cell_text = cell.text[:100]
                        table_html += f'<td style="border: 1px solid #4d4d4d; padding: 5px;">{cell_text}</td>'
                    table_html += '</tr>'
                table_html += '</table><br>'
                tables_html.append(table_html)

            # 暗黑模式完整HTML
            full_html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', '宋体', Arial, sans-serif;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        padding: 20px;
                        margin: 0;
                        line-height: 1.4;
                        background-color: #1e1e1e;
                        color: #e0e0e0;
                    }}
                    .info-bar {{
                        background-color: #2d2d2d;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #aaa;
                        border: 1px solid #3d3d3d;
                    }}
                    h4 {{
                        margin: 15px 0 5px 0;
                        color: #f0f0f0;
                    }}
                    table {{
                        margin: 10px 0;
                        background-color: #252525;
                    }}
                    td {{
                        vertical-align: top;
                        border: 1px solid #4d4d4d;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📄 文档共 {len(doc.paragraphs)} 段，{len(doc.tables)} 个表格
                    （预览前50段，前3个表格）
                    💡 提示：使用上方工具栏的 + / - 按钮可缩放预览
                </div>
                {''.join(html_parts) if html_parts else '<p style="color:gray;">无文字内容可预览</p>'}
                {''.join(tables_html) if tables_html else ''}
            </body>
            </html>
            """

            self.ui.web_template_preview.setHtml(full_html)

        except Exception as e:
            self.ui.web_template_preview.setHtml(
                f"<p style='color:#ff6666;padding:20px;'>❌ Word预览失败：{str(e)}<br><br>请确保已安装 python-docx</p>"
            )

    def _preview_excel_backup(self, file_path):
        """预览 Excel 文档（支持合并单元格）"""
        try:
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            # 加载工作簿
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active

            # 获取合并单元格区域
            merged_ranges = list(ws.merged_cells.ranges)

            # 获取数据范围
            min_row = ws.min_row or 1
            max_row = min(ws.max_row, 50)  # 最多50行
            min_col = ws.min_column or 1
            max_col = min(ws.max_column, 20)  # 最多20列

            # 构建数据矩阵
            data = []
            for row in range(min_row, max_row + 1):
                row_data = []
                for col in range(min_col, max_col + 1):
                    cell = ws.cell(row, col)
                    value = cell.value

                    # 检查是否在合并单元格中（且不是左上角）
                    is_merged = False
                    for merged_range in merged_ranges:
                        if (merged_range.min_row <= row <= merged_range.max_row and
                                merged_range.min_col <= col <= merged_range.max_col):
                            if row == merged_range.min_row and col == merged_range.min_col:
                                # 左上角单元格，正常显示
                                break
                            else:
                                # 被合并的单元格，置空
                                is_merged = True
                                break

                    if is_merged:
                        row_data.append('')
                    else:
                        # 处理空值
                        if value is None:
                            row_data.append('')
                        elif isinstance(value, (int, float)):
                            # 保留两位小数
                            row_data.append(f'{value:.2f}' if value % 1 != 0 else str(value))
                        else:
                            row_data.append(str(value)[:100])
                data.append(row_data)

            # 获取列宽
            col_widths = {}
            for col in range(min_col, max_col + 1):
                col_letter = get_column_letter(col)
                width = ws.column_dimensions[col_letter].width
                if width:
                    col_widths[col] = min(width, 200)
                else:
                    col_widths[col] = 100

            # 构建表格HTML
            html_rows = []

            # 表头（列号）
            html_rows.append('<thead><tr>')
            for col in range(min_col, max_col + 1):
                col_letter = get_column_letter(col)
                width = col_widths.get(col, 100)
                html_rows.append(f'<th style="min-width: {width}px;">{col_letter}</th>')
            html_rows.append('</thead>')

            # 表体
            html_rows.append('<tbody>')
            for row_idx, row_data in enumerate(data):
                html_rows.append('<tr>')
                for value in row_data:
                    html_rows.append(f'<td>{value}</td>')
                html_rows.append('</tr>')
            html_rows.append('</tbody>')

            # 处理合并单元格的HTML
            # 需要生成 col/rowspan 属性
            # 这里简化处理，用CSS模拟合并效果

            # 更精确的方法：构建带合并信息的HTML
            html_table = self._build_excel_html_with_merges(ws, min_row, max_row, min_col, max_col)

            if html_table:
                final_html = html_table
            else:
                # 降级方案
                final_html = f"""
                <table class="excel-table">
                    {''.join(html_rows)}
                </table>
                """

            # 完整HTML
            full_html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', 'Segoe UI', Arial, sans-serif;
                        padding: 10px;
                        margin: 0;
                        background-color: #fff;
                    }}
                    .excel-table {{
                        border-collapse: collapse;
                        width: 100%;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                    }}
                    .excel-table th {{
                        background-color: #f2f2f2;
                        border: 1px solid #ddd;
                        padding: 8px;
                        text-align: center;
                        font-weight: bold;
                        position: sticky;
                        top: 0;
                    }}
                    .excel-table td {{
                        border: 1px solid #ddd;
                        padding: 6px 8px;
                        text-align: left;
                        white-space: nowrap;
                    }}
                    .excel-table tr:nth-child(even) {{
                        background-color: #f9f9f9;
                    }}
                    .excel-table tr:hover {{
                        background-color: #e8f4f8;
                    }}
                    .info-bar {{
                        background-color: #f5f5f5;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #666;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📊 工作表：{ws.title}，共 {ws.max_row} 行，{ws.max_column} 列
                    （预览区域：{min_row}-{max_row}行，{min_col}-{max_col}列）
                    合并单元格：{len(merged_ranges)} 处
                    💡 提示：使用上方工具栏的 + / - 按钮可缩放预览
                </div>
                {final_html}
            </body>
            </html>
            """

            self.ui.web_template_preview.setHtml(full_html)

        except Exception as e:
            self.ui.web_template_preview.setHtml(
                f"<p style='color:red;padding:20px;'>❌ Excel预览失败：{str(e)}<br><br>请确保已安装 openpyxl</p>"
            )

    def _preview_excel_dark(self, file_path):
        """预览 Excel 文档（支持合并单元格）- 暗黑模式"""
        try:
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            wb = load_workbook(file_path, data_only=True)
            ws = wb.active
            merged_ranges = list(ws.merged_cells.ranges)

            min_row = ws.min_row or 1
            max_row = min(ws.max_row, 50)
            min_col = ws.min_column or 1
            max_col = min(ws.max_column, 20)

            # 构建带合并信息的HTML
            html_table = self._build_excel_html_with_merges(ws, min_row, max_row, min_col, max_col)

            if not html_table:
                html_table = "<p>无法生成表格预览</p>"

            # 暗黑模式完整HTML
            full_html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', 'Segoe UI', Arial, sans-serif;
                        padding: 10px;
                        margin: 0;
                        background-color: #1e1e1e;
                        color: #e0e0e0;
                    }}
                    .excel-table {{
                        border-collapse: collapse;
                        width: 100%;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        background-color: #252525;
                    }}
                    .excel-table th {{
                        background-color: #3d3d3d;
                        border: 1px solid #4d4d4d;
                        padding: 8px;
                        text-align: center;
                        font-weight: bold;
                        color: #f0f0f0;
                        position: sticky;
                        top: 0;
                    }}
                    .excel-table td {{
                        border: 1px solid #4d4d4d;
                        padding: 6px 8px;
                        text-align: left;
                        white-space: nowrap;
                        color: #d0d0d0;
                    }}
                    .excel-table tr:nth-child(even) {{
                        background-color: #2a2a2a;
                    }}
                    .excel-table tr:hover {{
                        background-color: #3a3a3a;
                    }}
                    .info-bar {{
                        background-color: #2d2d2d;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #aaa;
                        border: 1px solid #3d3d3d;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📊 工作表：{ws.title}，共 {ws.max_row} 行，{ws.max_column} 列
                    （预览区域：{min_row}-{max_row}行，{min_col}-{max_col}列）
                    合并单元格：{len(merged_ranges)} 处
                    💡 提示：使用上方工具栏的 + / - 按钮可缩放预览
                </div>
                {html_table}
            </body>
            </html>
            """

            self.ui.web_template_preview.setHtml(full_html)

        except Exception as e:
            self.ui.web_template_preview.setHtml(
                f"<p style='color:#ff6666;padding:20px;'>❌ Excel预览失败：{str(e)}<br><br>请确保已安装 openpyxl</p>"
            )

    def _build_excel_html_with_merges(self, ws, min_row, max_row, min_col, max_col):
        """构建带合并单元格的HTML表格（支持换行和对齐）"""
        try:
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import Alignment

            # 创建单元格矩阵
            cell_matrix = []
            for row in range(min_row, max_row + 1):
                row_cells = []
                for col in range(min_col, max_col + 1):
                    cell = ws.cell(row, col)
                    value = cell.value

                    # 获取对齐方式
                    align = cell.alignment
                    text_align = "left"
                    vertical_align = "top"

                    if align:
                        if align.horizontal:
                            hor_map = {
                                'center': 'center',
                                'centerContinuous': 'center',
                                'right': 'right',
                                'left': 'left',
                                'general': 'left'
                            }
                            text_align = hor_map.get(align.horizontal, 'left')
                        if align.vertical:
                            ver_map = {
                                'center': 'center',
                                'top': 'top',
                                'bottom': 'bottom'
                            }
                            vertical_align = ver_map.get(align.vertical, 'top')

                    # 处理换行符
                    if value and isinstance(value, str):
                        # 替换 \n 为 <br>
                        value = value.replace('\n', '<br>')
                    elif value is None:
                        value = ''
                    elif isinstance(value, (int, float)):
                        # 保留数字格式
                        value = f'{value:.2f}' if isinstance(value, float) and value % 1 != 0 else str(value)
                    else:
                        value = str(value)[:500]  # 限制长度

                    row_cells.append({
                        'value': value,
                        'row': row,
                        'col': col,
                        'merged': False,
                        'rowspan': 1,
                        'colspan': 1,
                        'align': text_align,
                        'valign': vertical_align
                    })
                cell_matrix.append(row_cells)

            # 标记合并单元格
            for merged in ws.merged_cells.ranges:
                if (merged.min_row > max_row or merged.max_row < min_row or
                        merged.min_col > max_col or merged.max_col < min_col):
                    continue

                start_row = max(merged.min_row, min_row) - min_row
                end_row = min(merged.max_row, max_row) - min_row
                start_col = max(merged.min_col, min_col) - min_col
                end_col = min(merged.max_col, max_col) - min_col

                if start_row < 0 or end_row >= len(cell_matrix) or start_col < 0 or end_col >= len(cell_matrix[0]):
                    continue

                # 设置左上角单元格
                cell_matrix[start_row][start_col]['rowspan'] = end_row - start_row + 1
                cell_matrix[start_row][start_col]['colspan'] = end_col - start_col + 1

                # 标记其他单元格
                for r in range(start_row, end_row + 1):
                    for c in range(start_col, end_col + 1):
                        if r == start_row and c == start_col:
                            continue
                        cell_matrix[r][c]['merged'] = True

            # 生成HTML
            html = ['<table class="excel-table">']

            # 表头
            html.append('<thead><tr>')
            for col in range(min_col, max_col + 1):
                col_letter = get_column_letter(col)
                # 获取列宽
                width = ws.column_dimensions[col_letter].width
                if width:
                    width_style = f'style="min-width: {min(width, 200)}px;"'
                else:
                    width_style = ''
                html.append(f'<th {width_style}>{col_letter}</th>')
            html.append('</tr></thead>')

            # 表体
            html.append('<tbody>')
            for row_cells in cell_matrix:
                html.append('<tr class="excel-row">')
                for cell in row_cells:
                    if cell['merged']:
                        continue

                    rowspan = cell['rowspan']
                    colspan = cell['colspan']
                    value = cell['value']
                    align = cell['align']
                    valign = cell['valign']

                    # 构建样式
                    style = f'text-align: {align}; vertical-align: {valign};'
                    if '\n' in value or '<br>' in value:
                        style += ' white-space: pre-wrap;'  # 允许换行

                    attrs = [f'style="{style}"']
                    if rowspan > 1:
                        attrs.append(f'rowspan="{rowspan}"')
                    if colspan > 1:
                        attrs.append(f'colspan="{colspan}"')

                    attrs_str = ' '.join(attrs)
                    html.append(f'<td {attrs_str}>{value}</td>')
                html.append('</tr>')
            html.append('</tbody>')

            html.append('</table>')
            return ''.join(html)

        except Exception as e:
            print(f"构建合并单元格HTML失败：{e}")
            return None

    def clear_prompt(self):
        """清除输入框内容"""
        self.ui.edit_prompt.clear()
        QMessageBox.information(self, "提示", "输入框已清空")

    def generate_document(self):
        """生成文档"""
        prompt = self.ui.edit_prompt.toPlainText().strip()
        if not prompt:
            QMessageBox.warning(self, "提示", "请输入需求描述")
            return

        format_text = self.ui.combo_format.currentText()
        format_map = {
            "Word (.docx)": "docx",
            "Excel (.xlsx)": "xlsx",
            "PPT (.pptx)": "pptx"
        }
        file_ext = format_map.get(format_text, "docx")

        # 获取选中的模板（可选）
        template_item = self.ui.list_templates.currentItem()
        template_path = None
        if template_item:
            template_name = template_item.text()
            template_path = os.path.join(self.template_dir, template_name)
            if init_config.isInTestMode:
                print(f"✅ 已选中模板：{template_name}")
        else:
            if init_config.isInTestMode:
                print(f"❌ 未选中任何模板，将使用无模板模式")

        # 禁用按钮，显示进度（暗黑模式适配）
        self.ui.btn_generate.setEnabled(False)

        if self.isInDarkMode:
            wait_html = """
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {
                        background-color: #1e1e1e;
                        color: #e0e0e0;
                        font-family: '微软雅黑', Arial, sans-serif;
                        display: flex;
                        justify-content: center;
                        align-items: center;
                        height: 100vh;
                        margin: 0;
                    }
                    .wait-container {
                        text-align: center;
                    }
                    .spinner {
                        width: 40px;
                        height: 40px;
                        border: 3px solid #3d3d3d;
                        border-top-color: #aaa;
                        border-radius: 50%;
                        animation: spin 1s linear infinite;
                        margin: 0 auto 15px auto;
                    }
                    @keyframes spin {
                        to { transform: rotate(360deg); }
                    }
                    .wait-text {
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #aaa;
                    }
                </style>
            </head>
            <body>
                <div class="wait-container">
                    <div class="spinner"></div>
                    <div class="wait-text">⏳ 正在生成文档，请稍候...</div>
                </div>
            </body>
            </html>
            """
        else:
            wait_html = "<p style='color:gray;text-align:center;padding:20px'>⏳ 正在生成文档，请稍候...</p>"

        self.ui.web_output_preview.setHtml(wait_html)

        # 在新线程中执行
        self.worker = DocumentGeneratorWorker(prompt, file_ext, template_path)
        self.worker.finished.connect(self._on_generate_finished)
        self.worker.error.connect(self._on_generate_error)
        self.worker.start()

    def _on_generate_finished(self, output_path):
        """生成完成"""
        self.ui.btn_generate.setEnabled(True)

        # 保存当前文件路径（供右键菜单使用）
        self.current_output_file = output_path

        # 先清空右侧预览区
        self.ui.web_output_preview.setHtml("")

        # 根据文件类型调用对应的输出预览方法
        if output_path.endswith('.docx'):
            if self.isInDarkMode:
                self._preview_output_word_dark(output_path)  # 如果还没有，可以类似添加
            else:
                self._preview_output_excel(output_path)
        elif output_path.endswith('.xlsx'):
            if self.isInDarkMode:
                self._preview_output_excel_dark(output_path)  # ← 使用新方法
            else:
                self._preview_output_excel(output_path)
        else:
            self.ui.web_output_preview.setHtml(
                f"<p>✅ 文档已生成</p><p>保存路径：{output_path}</p>"
            )

        # 弹出成功提示
        QMessageBox.information(self, "成功", f"文档已生成：{os.path.basename(output_path)}")

    def _preview_output_excel(self, file_path):
        """在输出预览区显示 Excel 文件内容"""
        try:
            import pandas as pd
            import os

            # 读取 Excel 文件（前50行）
            df = pd.read_excel(file_path, nrows=50)

            # 获取文件名
            file_name = os.path.basename(file_path)

            html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', 'Segoe UI', Arial, sans-serif;
                        padding: 10px;
                        margin: 0;
                        background: #fff;
                    }}
                    .info-bar {{
                        background: #f5f5f5;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #666;
                    }}
                    table {{
                        border-collapse: collapse;
                        width: 100%;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                    }}
                    th {{
                        background-color: #f2f2f2;
                        border: 1px solid #ddd;
                        padding: 8px;
                        text-align: center;
                        font-weight: bold;
                        position: sticky;
                        top: 0;
                    }}
                    td {{
                        border: 1px solid #ddd;
                        padding: 6px 8px;
                    }}
                    tr:nth-child(even) {{
                        background-color: #f9f9f9;
                    }}
                    tr:hover {{
                        background-color: #e8f4f8;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📊 生成文件：{file_name}，共 {len(df)} 行，{len(df.columns)} 列
                    💡 提示：使用上方工具栏的 +/- 按钮可缩放预览
                </div>
                {df.to_html()}
            </body>
            </html>
            """

            self.ui.web_output_preview.setHtml(html)

        except Exception as e:
            self.ui.web_output_preview.setHtml(
                f"<p style='color:red;padding:20px;'>❌ 预览失败：{str(e)}</p>"
            )

    def _preview_output_word(self, file_path):
        """在输出预览区显示 Word 文件内容"""
        try:
            from docx import Document

            doc = Document(file_path)
            file_name = os.path.basename(file_path)

            # 提取段落文本
            paragraphs = []
            for para in doc.paragraphs[:50]:
                if para.text.strip():
                    paragraphs.append(para.text)

            html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', 'Segoe UI', Arial, sans-serif;
                        padding: 20px;
                        background: #fff;
                        line-height: 1.4;
                    }}
                    .info-bar {{
                        background: #f5f5f5;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #666;
                    }}
                    .content {{
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📄 生成文件：{file_name}，共 {len(doc.paragraphs)} 段
                    💡 提示：使用上方工具栏的 +/- 按钮可缩放预览
                </div>
                <div class="content">
                    {'<p>'.join(paragraphs) if paragraphs else '<p>无文字内容</p>'}
                </div>
            </body>
            </html>
            """

            self.ui.web_output_preview.setHtml(html)

        except Exception as e:
            self.ui.web_output_preview.setHtml(
                f"<p style='color:red'>预览失败：{str(e)}</p>"
            )

    def _preview_output_excel_dark(self, file_path):
        """在输出预览区显示 Excel 文件内容（暗黑模式）"""
        try:
            import pandas as pd
            import os

            # 读取 Excel 文件（前50行）
            df = pd.read_excel(file_path, nrows=50)

            # 获取文件名
            file_name = os.path.basename(file_path)

            # 暗黑模式 CSS
            html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', 'Segoe UI', Arial, sans-serif;
                        padding: 10px;
                        margin: 0;
                        background-color: #1e1e1e;
                        color: #e0e0e0;
                    }}
                    .info-bar {{
                        background-color: #2d2d2d;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #aaa;
                        border: 1px solid #3d3d3d;
                    }}
                    table {{
                        border-collapse: collapse;
                        width: 100%;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        background-color: #252525;
                    }}
                    th {{
                        background-color: #3d3d3d;
                        border: 1px solid #4d4d4d;
                        padding: 8px;
                        text-align: center;
                        font-weight: bold;
                        color: #f0f0f0;
                        position: sticky;
                        top: 0;
                    }}
                    td {{
                        border: 1px solid #4d4d4d;
                        padding: 6px 8px;
                        color: #d0d0d0;
                    }}
                    tr:nth-child(even) {{
                        background-color: #2a2a2a;
                    }}
                    tr:hover {{
                        background-color: #3a3a3a;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📊 生成文件：{file_name}，共 {len(df)} 行，{len(df.columns)} 列
                    💡 提示：使用上方工具栏的 +/- 按钮可缩放预览
                </div>
                {df.to_html()}
            </body>
            </html>
            """

            self.ui.web_output_preview.setHtml(html)

        except Exception as e:
            self.ui.web_output_preview.setHtml(
                f"<p style='color:#ff6666;padding:20px;'>❌ 预览失败：{str(e)}</p>"
            )

    def _preview_output_word_dark(self, file_path):
        """在输出预览区显示 Word 文件内容（暗黑模式）"""
        try:
            from docx import Document
            import os

            doc = Document(file_path)
            file_name = os.path.basename(file_path)

            # 提取段落文本
            paragraphs = []
            for para in doc.paragraphs[:50]:
                if para.text.strip():
                    paragraphs.append(para.text)

            # 暗黑模式 CSS
            html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <style>
                    body {{
                        font-family: '微软雅黑', 'Segoe UI', Arial, sans-serif;
                        padding: 20px;
                        margin: 0;
                        background-color: #1e1e1e;
                        color: #e0e0e0;
                        line-height: 1.5;
                    }}
                    .info-bar {{
                        background-color: #2d2d2d;
                        padding: 8px;
                        margin-bottom: 10px;
                        border-radius: 4px;
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        color: #aaa;
                        border: 1px solid #3d3d3d;
                    }}
                    .content {{
                        font-size: {DEFAULT_FONT_SIZE-4}pt;
                        background-color: #252525;
                        padding: 15px;
                        border-radius: 4px;
                        border: 1px solid #3d3d3d;
                    }}
                </style>
            </head>
            <body>
                <div class="info-bar">
                    📄 生成文件：{file_name}，共 {len(doc.paragraphs)} 段
                    💡 提示：使用上方工具栏的 +/- 按钮可缩放预览
                </div>
                <div class="content">
                    {'<p>'.join(paragraphs) if paragraphs else '<p>无文字内容</p>'}
                </div>
            </body>
            </html>
            """

            self.ui.web_output_preview.setHtml(html)

        except Exception as e:
            self.ui.web_output_preview.setHtml(
                f"<p style='color:#ff6666'>预览失败：{str(e)}</p>"
            )

    def _on_generate_error(self, error_msg):
        """生成失败"""
        self.ui.btn_generate.setEnabled(True)
        self.ui.web_output_preview.setHtml(
            f"<p style='color:red'>❌ 生成失败：{error_msg}</p>"
        )
        QMessageBox.critical(self, "错误", error_msg)

    def open_settings(self):
        dialog = DocsSettingDialog(self)
        if dialog.exec():
            # 可选：重新加载设置，更新当前窗口的配置
            settings = dialog.get_settings()
            # 比如更新输出目录
            self.output_dir = settings.get("output_dir", "./outputs")
            QMessageBox.information(self, "提示", "设置已保存")

    def _setup_preview_context_menu(self):
        """为预览区添加右键菜单"""
        self.ui.web_output_preview.setContextMenuPolicy(Qt.CustomContextMenu)
        self.ui.web_output_preview.customContextMenuRequested.connect(self._show_preview_menu)

    def _show_preview_menu(self, position):
        """显示右键菜单"""
        from PySide6.QtWidgets import QMenu

        menu = QMenu()

        # 获取当前预览的文件路径（需要在生成时保存）
        current_file = getattr(self, 'current_output_file', None)

        if current_file and os.path.exists(current_file):
            open_action = menu.addAction("📂 打开文件")
            open_action.triggered.connect(lambda: self._open_file(current_file))

            reveal_action = menu.addAction("📁 在文件夹中显示")
            reveal_action.triggered.connect(lambda: self._reveal_file(current_file))

            menu.addSeparator()

            copy_action = menu.addAction("📋 复制文件路径")
            copy_action.triggered.connect(lambda: self._copy_file_path(current_file))

        menu.exec_(self.ui.web_output_preview.mapToGlobal(position))

    def _open_file(self, file_path):
        """用默认程序打开文件"""
        import subprocess
        import sys
        if sys.platform == 'win32':
            os.startfile(file_path)
        elif sys.platform == 'darwin':  # macOS
            subprocess.run(['open', file_path])
        else:  # Linux
            subprocess.run(['xdg-open', file_path])

    def _reveal_file(self, file_path):
        """在文件管理器中显示"""
        import subprocess
        import sys
        dir_path = os.path.dirname(file_path)
        if sys.platform == 'win32':
            subprocess.run(['explorer', dir_path])
        elif sys.platform == 'darwin':
            subprocess.run(['open', dir_path])
        else:
            subprocess.run(['xdg-open', dir_path])

    def _copy_file_path(self, file_path):
        """复制文件路径到剪贴板"""
        from PySide6.QtWidgets import QApplication
        QApplication.clipboard().setText(file_path)


class DocumentGeneratorWorker(QThread):
    finished = Signal(str)  # 输出文件路径
    error = Signal(str)  # 错误信息

    def __init__(self, prompt, file_ext, template_path=None):
        super().__init__()
        self.prompt = prompt
        self.file_ext = file_ext
        self.template_path = template_path

    def run(self):
        """主入口：协调整个生成流程"""
        try:
            # 1. 加载配置
            config = self._load_config()
            if not config:
                return

            # 2. 调用 API 生成内容
            content = self._call_api(config)
            if content is None:
                return

            # 3. 保存为文件
            output_path = self._save_to_file(content, config)

            # 4. 完成
            self.finished.emit(output_path)

        except Exception as e:
            self.error.emit(str(e))

    # ==================== 步骤1：加载配置 ====================

    def _load_config(self):
        """加载文档助手配置"""
        from init_config import load_doc_assistant_config

        config = load_doc_assistant_config()
        api_key = config.get("api_key", "")

        if not api_key:
            self.error.emit("请先在设置中填写 API Key")
            return None

        # 确保输出目录存在
        output_dir = config.get("output_dir", "./outputs")
        os.makedirs(output_dir, exist_ok=True)

        return config

    # ==================== 步骤2：调用 API ====================

    def _call_api(self, config):
        """调用千问 API 生成文档内容"""
        from openai import OpenAI

        api_key = config.get("api_key", "")
        model = config.get("model", "qwen-max")
        temperature = config.get("temperature", 0.7)

        client = OpenAI(
            api_key=api_key,
            base_url="https://dashscope.aliyuncs.com/compatible-mode/v1"
        )

        # 根据格式构建系统提示词
        system_prompt = self._build_system_prompt()

        try:
            response = client.chat.completions.create(
                model=model,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": self._build_user_prompt()}
                ],
                temperature=temperature
            )
            content = response.choices[0].message.content

            if init_config.isInTestMode:
                # ===== 添加调试打印 =====
                print("\n" + "=" * 60)
                print("🔍 千问 API 原始返回内容：")
                print("=" * 60)
                print(content)
                print("=" * 60)
                print(f"📊 内容长度：{len(content)} 字符")
                print(f"📝 前200字符：{content[:200]}")
                print("=" * 60 + "\n")

            return content

        except Exception as e:
            self.error.emit(f"API 调用失败：{str(e)}")
            return None

    def _build_system_prompt(self):
        """根据文件格式构建系统提示词"""
        base_prompt = "你是一个专业的文档生成助手。"

        if self.file_ext == 'xlsx':
            return base_prompt + """
                                    请根据用户的需求，生成一份 Excel 表格数据。

                                    【重要格式要求】
                                        - 只输出一个 JSON 二维数组，不要输出任何其他文字
                                """
        elif self.file_ext == 'docx':
            return base_prompt + """
                                    请生成一份 Word 文档，直接输出文档正文内容，支持 Markdown 格式。
                                    可以使用标题、列表、表格等元素，确保内容结构清晰、专业。
                                """
        elif self.file_ext == 'pptx':
            return base_prompt + """
                                    请生成一份 PPT 演示文稿内容，用 --- 分隔每一张幻灯片。
                                    每张幻灯片包含标题和内容，可以使用 Markdown 格式。
                                    """
        else:
            return base_prompt + "请直接输出文档内容。"

    def _build_user_prompt(self):
        prompt = f"请生成一份 {self.file_ext.upper()} 文档：{self.prompt}"

        if self.template_path and os.path.exists(self.template_path):
            template_content = self._read_template_content()
            if template_content and template_content.strip():  # 加这个判断
                prompt += f"\n\n请参考以下模板..."
        else:
            prompt += "\n\n请直接根据需求生成文档，不需要参考任何模板。"

        return prompt

    def _read_template_content(self):
        """读取模板内容（根据文件类型）"""
        try:
            ext = os.path.splitext(self.template_path)[1].lower()

            if ext == '.xlsx':
                import pandas as pd
                df = pd.read_excel(self.template_path, nrows=5)
                columns = df.columns.tolist()
                sample = df.head(2).to_string()
                return f"列结构：{columns}\n示例数据：\n{sample}"

            elif ext == '.docx':
                from docx import Document
                doc = Document(self.template_path)
                paragraphs = []
                for para in doc.paragraphs[:30]:
                    if para.text.strip():
                        paragraphs.append(para.text)
                # 也读取表格结构
                tables_info = []
                for table in doc.tables[:2]:
                    headers = []
                    for cell in table.rows[0].cells:
                        headers.append(cell.text.strip())
                    tables_info.append(f"表格列：{', '.join(headers)}")

                content = '\n'.join(paragraphs)
                if tables_info:
                    content += '\n' + '\n'.join(tables_info)
                return content[:3000]  # 限制长度，避免超出 token 限制

            elif ext == '.pptx':
                # PPT 模板暂不支持读取
                return None

            else:
                return None

        except Exception as e:
            print(f"读取模板失败：{e}")
            return None

    # ==================== 步骤3：保存文件 ====================

    def _save_to_file(self, content, config):
        """将内容保存为文件"""
        output_dir = config.get("output_dir", "./outputs")
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filename = f"generated_{timestamp}.{self.file_ext}"
        output_path = os.path.join(output_dir, filename)

        if self.file_ext == 'docx':
            self._save_as_docx(content, output_path)
        elif self.file_ext == 'xlsx':
            self._save_as_excel(content, output_path)
        else:
            self._save_as_text(content, output_path)

        return output_path

    def _save_as_docx(self, content, output_path):
        """保存为 Word 文档"""
        from docx import Document
        from docx.shared import Inches

        doc = Document()
        doc.add_heading('生成的文档', 0)

        for line in content.split('\n'):
            if line.strip():
                # 简单处理 Markdown 标题
                if line.startswith('# '):
                    doc.add_heading(line[2:], 1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], 2)
                else:
                    doc.add_paragraph(line)

        doc.save(output_path)

    def _save_as_excel(self, content, output_path):
        """保存为 Excel 文档（智能解析）"""
        import pandas as pd
        import json
        import re

        # 尝试解析 JSON 二维数组
        if self._parse_json_to_excel(content, output_path):
            return

        # 尝试解析 Markdown 表格
        if self._parse_markdown_table_to_excel(content, output_path):
            return

        # 降级：保存为纯文本并提示
        self._save_as_text(content, output_path, warning=True)
        raise ValueError("API返回的内容无法解析为表格格式，已保存为纯文本")

    def _parse_json_to_excel(self, content, output_path):
        """解析 JSON 二维数组并保存为 Excel"""
        import pandas as pd
        import json
        import re

        try:
            # 匹配 JSON 数组
            json_match = re.search(r'\[\s*\[.*\]\s*\]', content, re.DOTALL)
            if not json_match:
                return False

            data = json.loads(json_match.group())
            if not isinstance(data, list) or len(data) == 0:
                return False

            # 第一行作为表头，其余作为数据
            df = pd.DataFrame(data[1:], columns=data[0])
            df.to_excel(output_path, index=False)
            return True

        except Exception as e:
            print(f"JSON 解析失败：{e}")
            return False

    def _parse_markdown_table_to_excel(self, content, output_path):
        """解析 Markdown 表格并保存为 Excel"""
        import pandas as pd

        try:
            lines = content.strip().split('\n')
            table_lines = []
            in_table = False

            for line in lines:
                if line.startswith('|') and '|' in line:
                    table_lines.append(line)
                    in_table = True
                elif in_table and line.strip() == '':
                    break
                elif in_table and not line.startswith('|'):
                    break

            if len(table_lines) < 2:
                return False

            # 解析表头
            headers = [h.strip() for h in table_lines[0].split('|')[1:-1]]

            # 解析数据行（跳过分隔行）
            data_rows = []
            for line in table_lines[2:]:
                row = [cell.strip() for cell in line.split('|')[1:-1]]
                if row:
                    data_rows.append(row)

            if not data_rows:
                return False

            df = pd.DataFrame(data_rows, columns=headers)
            df.to_excel(output_path, index=False)
            return True

        except Exception as e:
            print(f"Markdown 表格解析失败：{e}")
            return False

    def _save_as_text(self, content, output_path, warning=False):
        """保存为纯文本文件"""
        with open(output_path, 'w', encoding='utf-8') as f:
            if warning:
                f.write("⚠️ 内容无法转换为目标格式，以下是原始内容：\n\n")
            f.write(content)
